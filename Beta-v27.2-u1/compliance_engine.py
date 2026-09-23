"""Deterministic compliance evaluation for Sentinel X.

The engine deliberately keeps PASS/FAIL decisions out of AI providers. Rules
are stored as data, evaluate normalized telemetry, and preserve source
snapshots as evidence that can survive OpenSearch retention.
"""
import hashlib
import io
import json
import logging
import re
from datetime import datetime, timedelta

from openpyxl import load_workbook

from models import (
    ComplianceAuditLog,
    ComplianceControl,
    ComplianceDomain,
    ComplianceEvidence,
    ComplianceFinding,
    ComplianceFramework,
    ComplianceRequirement,
    ComplianceRemediation,
    ComplianceResult,
    ComplianceRisk,
    ComplianceRule,
    ComplianceRuleControl,
    ComplianceWorkbook,
    ComplianceWorkbookRecord,
    db,
)

logger = logging.getLogger(__name__)


def ensure_compliance_schema():
    """Apply the small additive changes needed by existing installations."""
    from sqlalchemy import text
    try:
        with db.engine.begin() as connection:
            connection.execute(text(
                'ALTER TABLE compliance_evidence '
                'ADD COLUMN IF NOT EXISTS control_id INTEGER '
                'REFERENCES compliance_control(id) ON DELETE SET NULL'
            ))
            connection.execute(text(
                'ALTER TABLE compliance_evidence ALTER COLUMN finding_id DROP NOT NULL'
            ))
    except Exception:
        # Fresh databases already receive the nullable model definition from
        # create_all; older engines may not support PostgreSQL ALTER syntax.
        logger.debug('Compliance additive schema migration was not required', exc_info=True)


def _value_at(source, path):
    value = source
    for part in str(path or '').split('.'):
        if not isinstance(value, dict):
            return None
        value = value.get(part)
    return value


def _compare(actual, operator, expected):
    operator = (operator or 'equals').lower()
    if operator in ('equals', '=='):
        return actual == expected or str(actual).lower() == str(expected).lower()
    if operator in ('not_equals', '!='):
        return not _compare(actual, 'equals', expected)
    if operator == 'contains':
        if isinstance(actual, (list, tuple, set)):
            return expected in actual or str(expected).lower() in [str(v).lower() for v in actual]
        return str(expected).lower() in str(actual or '').lower()
    if operator == 'regex':
        return bool(re.search(str(expected), str(actual or ''), re.IGNORECASE))
    try:
        left, right = float(actual), float(expected)
        if operator in ('greater_than', '>'):
            return left > right
        if operator in ('greater_than_equal', '>='):
            return left >= right
        if operator in ('less_than', '<'):
            return left < right
        if operator in ('less_than_equal', '<='):
            return left <= right
    except (TypeError, ValueError):
        return False
    return False


def evaluate_condition(source, condition):
    """Evaluate a nested AND/OR condition using deterministic operators."""
    if not condition:
        return False
    if 'all' in condition:
        return all(evaluate_condition(source, item) for item in condition['all'])
    if 'any' in condition:
        return any(evaluate_condition(source, item) for item in condition['any'])
    return _compare(
        _value_at(source, condition.get('field')),
        condition.get('operator'),
        condition.get('value'),
    )


def audit(action, object_type, object_id=None, user_id=None, old=None, new=None):
    db.session.add(ComplianceAuditLog(
        user_id=user_id,
        action=action,
        object_type=object_type,
        object_id=str(object_id) if object_id is not None else None,
        old_value=json.dumps(old, default=str) if old is not None else None,
        new_value=json.dumps(new, default=str) if new is not None else None,
    ))


def _severity_score(severity):
    return {'critical': 25, 'high': 16, 'medium': 9, 'low': 4}.get(
        str(severity or '').lower(), 1
    )


def ensure_default_catalog():
    """Seed extensible starter frameworks, controls, and one safe rule."""
    if ComplianceFramework.query.count():
        # The starter rule must work with the alert stream already collected
        # by this installation. Administrators can change it afterward.
        starter = ComplianceRule.query.filter_by(
            name='High-severity Wazuh alert requires review'
        ).first()
        if starter and starter.get_condition() == {
            'field': 'rule.level',
            'operator': 'greater_than_equal',
            'value': 12,
        }:
            starter.set_condition({
                'field': 'rule.level',
                'operator': 'greater_than_equal',
                'value': 5,
            })
            starter.name = 'Wazuh alert requires compliance review'
            db.session.commit()
        return

    definitions = [
        ('SOC2', 'SOC 2', 'Trust Services Criteria for security monitoring.'),
        ('ISO27001', 'ISO/IEC 27001', 'Information security management controls.'),
        ('INTERNAL', 'Internal Security Policies', 'Organization-owned security requirements.'),
    ]
    controls = []
    for code, name, description in definitions:
        framework = ComplianceFramework(code=code, name=name, description=description, version='2026.1')
        domain = ComplianceDomain(code=f'{code}-DOMAIN', name='Security and monitoring')
        requirement = ComplianceRequirement(
            code=f'{code}-REQ-01',
            title='Security events are monitored and reviewed',
            description='Security telemetry is collected, monitored, and retained for review.',
        )
        control = ComplianceControl(
            code=f'{code}-CTRL-01',
            title='Security monitoring and alert review',
            description='High-severity security telemetry must be evaluated and tracked to resolution.',
            control_type='automated',
        )
        domain.requirements.append(requirement)
        requirement.controls.append(control)
        framework.domains.append(domain)
        db.session.add(framework)
        controls.append(control)

    db.session.flush()
    rule = ComplianceRule(
        name='Wazuh alert requires compliance review',
        source='wazuh',
        severity='high',
        schedule='hourly',
    )
    rule.set_condition({
        'field': 'rule.level',
        'operator': 'greater_than_equal',
        'value': 5,
    })
    db.session.add(rule)
    db.session.flush()
    for control in controls:
        db.session.add(ComplianceRuleControl(rule_id=rule.id, control_id=control.id))
    audit('catalog_seeded', 'framework', new={'frameworks': [item[0] for item in definitions]})
    db.session.commit()


def _cell(value):
    if hasattr(value, 'value'):
        value = value.value
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _row_values(row):
    return [_cell(value) for value in row]


def _is_control_code(value):
    return bool(re.match(r'^CC\d+\.\d+$', str(value or '').strip(), re.IGNORECASE))


def _framework_control(framework, code, title, domain_names):
    """Upsert the native SOC 2 hierarchy for a workbook control code."""
    section = str(code).upper().split('.')[0]
    domain_name = domain_names.get(section, section)
    domain = next((item for item in framework.domains if item.code == section), None)
    if not domain:
        domain = ComplianceDomain(code=section, name=domain_name)
        framework.domains.append(domain)
    elif domain_name and domain.name == section:
        domain.name = domain_name
    requirement = next((item for item in domain.requirements if item.code == section), None)
    if not requirement:
        requirement = ComplianceRequirement(code=section, title=domain_name)
        domain.requirements.append(requirement)
    control = next((item for item in requirement.controls if item.code == code), None)
    if not control:
        control = ComplianceControl(
            code=code,
            title=title or code,
            description=title or '',
            control_type='manual',
        )
        requirement.controls.append(control)
    elif title and (not control.title or control.title == control.code):
        control.title = title
        control.description = title
    return control


def _record(workbook_id, sheet_name, record_type, record_key, title, data, control=None):
    return ComplianceWorkbookRecord(
        workbook_id=workbook_id,
        control_id=control.id if control else None,
        sheet_name=sheet_name,
        record_type=record_type,
        record_key=str(record_key) if record_key is not None else None,
        title=str(title or record_key or sheet_name),
        status=str(data.get('Status') or data.get('status') or '') or None,
        data_json=json.dumps(data, default=str, ensure_ascii=False),
    )


def import_soc2_workbook(file_bytes, filename, user_id=None):
    """Import the workbook's sheets into native SOC 2 controls and drill-down records."""
    digest = hashlib.sha256(file_bytes).hexdigest()
    existing = ComplianceWorkbook.query.filter_by(file_sha256=digest).first()
    if existing:
        return {
            'workbook_id': existing.id,
            'duplicate': True,
            'sheets': sorted({record.sheet_name for record in existing.records}),
            'controls': len({record.control_id for record in existing.records if record.control_id}),
            'records': len(existing.records),
        }

    workbook_file = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    framework = ComplianceFramework.query.filter_by(code='SOC2').first()
    if not framework:
        framework = ComplianceFramework(
            code='SOC2',
            name='SOC 2',
            description='AICPA Trust Services Criteria',
            version='2017',
        )
        db.session.add(framework)
        db.session.flush()
    workbook = ComplianceWorkbook(
        framework_id=framework.id,
        name='SOC 2 organization workbook',
        original_filename=filename,
        file_sha256=digest,
        imported_by=user_id,
    )
    db.session.add(workbook)
    db.session.flush()

    all_rows = {}
    domain_names = {}
    controls = {}
    for sheet_name in workbook_file.sheetnames:
        rows = [_row_values(row) for row in workbook_file[sheet_name].iter_rows()]
        rows = [row for row in rows if any(value is not None for value in row)]
        all_rows[sheet_name] = rows
        for row in rows:
            first = str(row[0] or '').strip()
            match = re.match(r'^(CC\d+)\s*[—–-]\s*(.+)$', first)
            if match:
                domain_names[match.group(1).upper()] = first

    actual_name = next((name for name in all_rows if name.lower().startswith('actual control')), None)
    if actual_name:
        current_domain = None
        for row in all_rows[actual_name]:
            first = str(row[0] or '').strip()
            match = re.match(r'^(CC\d+)\s*[—–-]\s*(.+)$', first)
            if match:
                current_domain = match.group(1).upper()
                domain_names[current_domain] = first
            elif _is_control_code(first):
                control = _framework_control(framework, first.upper(), row[1] if len(row) > 1 else first, domain_names)
                controls[first.upper()] = control
                db.session.flush()
                data = {'ID': first, 'Control Summary': row[1] if len(row) > 1 else None}
                db.session.add(_record(workbook.id, actual_name, 'control', first, row[1], data, control))

    for sheet_name, rows in all_rows.items():
        lowered = sheet_name.lower()
        if lowered.startswith('actual control'):
            continue
        header_index = None
        for index, row in enumerate(rows[:8]):
            if any(str(value or '').strip() in ('Control ID', 'Risk ID', 'Policy Name', 'SOC 2 Control ID') for value in row):
                header_index = index
                break
        if header_index is None:
            # Preserve overview/section rows even when a sheet has no table header.
            for row_number, row in enumerate(rows):
                data = {f'Column {i + 1}': value for i, value in enumerate(row) if value is not None}
                if data:
                    db.session.add(_record(workbook.id, sheet_name, 'overview', row_number + 1, str(row[0]), data))
            continue

        headers = [str(value).strip() if value is not None else f'Column {i + 1}'
                   for i, value in enumerate(rows[header_index])]
        record_type = (
            'implementation_plan' if 'implementation' in lowered else
            'tsc_mapping' if 'tsc' in lowered else
            'evidence_register' if 'evidence' in lowered else
            'policy_mapping' if 'policy' in lowered else
            'risk_register' if 'risk' in lowered else
            'ownership' if 'owner' in lowered or 'responsibility' in lowered else
            'register'
        )
        grouped = {}
        for row in rows[header_index + 1:]:
            values = list(row) + [None] * max(0, len(headers) - len(row))
            data = {headers[i]: values[i] for i in range(len(headers)) if values[i] is not None}
            if not data:
                continue
            key = data.get('Control ID') or data.get('SOC 2 Control ID') or data.get('Risk ID') or data.get('Policy Name')
            if record_type == 'implementation_plan':
                if _is_control_code(key):
                    grouped[str(key).upper()] = data
                    grouped[str(key).upper()]['Implementation Steps'] = [data.get('Implementation Steps')] if data.get('Implementation Steps') else []
                elif grouped and data.get('Implementation Steps'):
                    current = list(grouped)[-1]
                    grouped[current].setdefault('Implementation Steps', []).append(data['Implementation Steps'])
                continue
            title = data.get('Control Name') or data.get('Control Description') or data.get('Policy Name') or data.get('Risk Description') or data.get('Control Name / Area') or key
            control = controls.get(str(key).upper()) if _is_control_code(key) else None
            db.session.add(_record(workbook.id, sheet_name, record_type, key, title, data, control))
        for key, data in grouped.items():
            control = controls.get(key)
            if not control:
                control = _framework_control(framework, key, data.get('Control Description') or key, domain_names)
                controls[key] = control
                db.session.flush()
            db.session.add(_record(workbook.id, sheet_name, record_type, key, data.get('Control Description') or key, data, control))

    db.session.add(ComplianceAuditLog(
        user_id=user_id,
        action='soc2_workbook_imported',
        object_type='workbook',
        object_id=str(workbook.id),
        new_value=json.dumps({
            'filename': filename,
            'sheets': workbook_file.sheetnames,
            'controls': len(controls),
        }),
    ))
    db.session.commit()
    return {
        'workbook_id': workbook.id,
        'duplicate': False,
        'sheets': workbook_file.sheetnames,
        'controls': len(controls),
        'records': ComplianceWorkbookRecord.query.filter_by(workbook_id=workbook.id).count(),
    }


def _parse_timestamp(value):
    if not value:
        return datetime.utcnow()
    try:
        return datetime.fromisoformat(str(value).replace('Z', '+00:00')).replace(tzinfo=None)
    except (TypeError, ValueError):
        return datetime.utcnow()


def _extract_asset(source):
    agent = source.get('agent') if isinstance(source.get('agent'), dict) else {}
    host = source.get('host') if isinstance(source.get('host'), dict) else {}
    return agent.get('name') or host.get('name') or agent.get('id') or host.get('id') or 'Unknown asset'


def evaluate_wazuh(server_key='primary', hours=24, limit=200):
    """Evaluate recent Wazuh telemetry and upsert deduplicated findings."""
    from opensearch_api import OpenSearchAPI

    start = (datetime.utcnow() - timedelta(hours=max(1, min(hours, 720)))).isoformat()
    end = datetime.utcnow().isoformat()
    api = OpenSearchAPI(server_key=server_key)
    if not api.client:
        raise RuntimeError('OpenSearch is not available for this monitoring server')
    response = api.client.search(
        index=api.index_pattern,
        body={
            'size': max(1, min(limit, 1000)),
            'sort': [{'@timestamp': {'order': 'desc'}}],
            'query': {'range': {'@timestamp': {'gte': start, 'lte': end}}},
        },
    )
    hits = response.get('hits', {}).get('hits', [])
    rules = ComplianceRule.query.filter_by(enabled=True, source='wazuh').all()
    evaluated = findings_created = findings_updated = rule_matches = 0
    observed_levels = {}

    for hit in hits:
        source = hit.get('_source') or {}
        level = _value_at(source, 'rule.level')
        try:
            level_key = str(int(float(level)))
        except (TypeError, ValueError):
            level_key = str(level or 'unknown')
        observed_levels[level_key] = observed_levels.get(level_key, 0) + 1
        for rule in rules:
            if not evaluate_condition(source, rule.get_condition()):
                continue
            rule_matches += 1
            controls = [link.control for link in rule.controls if link.control]
            if not controls:
                continue
            asset = _extract_asset(source)
            timestamp = _parse_timestamp(source.get('@timestamp'))
            rule_info = source.get('rule') if isinstance(source.get('rule'), dict) else {}
            title = f"{rule.name}: {rule_info.get('description', 'telemetry matched')}"
            message = (
                f"Deterministic rule matched on {asset}. "
                f"Source rule {rule_info.get('id', 'unknown')}."
            )
            for control in controls:
                fingerprint = hashlib.sha256(
                    '|'.join([
                        server_key, asset, str(rule.id), str(control.id), 'wazuh',
                    ]).encode()
                ).hexdigest()
                finding = ComplianceFinding.query.filter_by(fingerprint=fingerprint).first()
                if finding:
                    finding.last_detected = timestamp
                    finding.occurrence_count = (finding.occurrence_count or 0) + 1
                    findings_updated += 1
                else:
                    finding = ComplianceFinding(
                        fingerprint=fingerprint,
                        title=title,
                        description=message,
                        source='wazuh',
                        server_key=server_key,
                        asset=asset,
                        severity=rule.severity,
                        risk_score=_severity_score(rule.severity),
                        rule_id=rule.id,
                        control_id=control.id,
                        first_detected=timestamp,
                        last_detected=timestamp,
                        remediation='Review the source alert, remediate the affected asset, and request verification.',
                    )
                    db.session.add(finding)
                    db.session.flush()
                    db.session.add(ComplianceRisk(
                        finding_id=finding.id,
                        name=title,
                        description=message,
                        likelihood=3,
                        impact=5 if rule.severity in ('high', 'critical') else 3,
                        inherent_risk=finding.risk_score,
                        residual_risk=finding.risk_score,
                    ))
                    findings_created += 1
                control.status = 'failed'
                db.session.add(ComplianceResult(
                    rule_id=rule.id,
                    control_id=control.id,
                    server_key=server_key,
                    asset=asset,
                    status='fail',
                    message=message,
                ))
                snapshot = {
                    'index': hit.get('_index'),
                    'document_id': hit.get('_id'),
                    'timestamp': source.get('@timestamp'),
                    'asset': asset,
                    'rule': rule_info,
                    'source': source,
                }
                encoded = json.dumps(snapshot, sort_keys=True, default=str)
                digest = hashlib.sha256(encoded.encode()).hexdigest()
                if not ComplianceEvidence.query.filter_by(finding_id=finding.id, sha256=digest).first():
                    db.session.add(ComplianceEvidence(
                        finding_id=finding.id,
                        source='wazuh',
                        server_key=server_key,
                        source_index=hit.get('_index'),
                        source_document_id=hit.get('_id'),
                        snapshot_json=encoded,
                        sha256=digest,
                    ))
            evaluated += 1
    db.session.commit()
    return {
        'evaluated_events': evaluated,
        'collected_events': len(hits),
        'rule_matches': rule_matches,
        'observed_levels': observed_levels,
        'findings_created': findings_created,
        'findings_updated': findings_updated,
        'server_key': server_key,
        'window_hours': hours,
    }